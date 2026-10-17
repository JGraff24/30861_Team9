"""Non-deployed preservation experiment: openpyxl round trips vs targeted OOXML edits.

Run from the repository root:
    uv run python -m experiments.preservation.run_experiment [--keep DIR]

Outputs are temporary copies (a temp directory unless --keep is given). The source
workbook is generated in memory and never modified. This is NOT production remediation.
"""

import argparse
import hashlib
import io
import json
import sys
import tempfile
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tests" / "unit"))

from ooxml_builder import (  # noqa: E402
    Drawing,
    Sheet,
    Workbook,
    anchor,
    chart,
    cnvpr,
    drawing_xml,
    pic,
    shape,
)

from experiments.preservation.snapshot import diff, snapshot  # noqa: E402

SHEET1 = "xl/worksheets/sheet1.xml"
DRAWING1 = "xl/drawings/drawing1.xml"
ALT_BEFORE = b'<xdr:cNvPr id="2" name="Picture 1"/>'
ALT_AFTER = b'<xdr:cNvPr id="2" name="Picture 1" descr="Photo of the sales team" title="Sales"/>'


def source_workbook() -> bytes:
    cells = (
        '<row r="1"><c r="A1" t="inlineStr"><is><t>Quarter</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>Total</t></is></c></row>'
        '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>0</v></c></row>'
        '<row r="3"><c r="A3"><v>2</v></c><c r="B3"><f>B2+10</f><v>10</v></c></row>'
        '<row r="4"><c r="A4" t="inlineStr"><is><t>Click here</t></is></c></row>'
    )
    links = (
        '<hyperlinks><hyperlink ref="A4" location="\'Expenses\'!A1" display="Click here"/>'
        "</hyperlinks>"
    )
    d1 = drawing_xml(
        anchor(pic(cnvpr("2", "Picture 1"))),
        anchor(chart(cnvpr("3", "Chart 1", title="Totals", descr="Totals rise to 10"))),
        anchor(shape(cnvpr("4", "Note", title="Note", descr="Excludes tax"), "Note")),
    )
    expenses = '<row r="1"><c r="A1"><f>Sheet1!B3*2</f><v>20</v></c></row>'
    wb = Workbook(
        [
            Sheet("Sheet1", SHEET1, "1", cells=cells, tail=links, drawing=Drawing(DRAWING1, d1)),
            Sheet("Expenses", "xl/worksheets/sheet2.xml", "2", cells=expenses),
        ]
    )
    parts = wb.parts()
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
        b"</sheets></workbook>",
        b'</sheets><definedNames><definedName name="Totals">Sheet1!$B$2:$B$3</definedName>'
        b"</definedNames></workbook>",
    )
    from ooxml_builder import zip_parts

    return zip_parts(parts)


# -- strategies --------------------------------------------------------------------------


def targeted(data: bytes, edits: dict[str, Callable[[bytes], bytes]]) -> bytes:
    """Copy every entry byte-for-byte, rewriting only the named parts."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename in edits:
                content = edits[info.filename](content)
            target.writestr(info, content)
    return out.getvalue()


def replace_once(old: bytes, new: bytes) -> Callable[[bytes], bytes]:
    def edit(content: bytes) -> bytes:
        if content.count(old) != 1:
            raise ValueError(f"expected exactly one {old!r}")
        return content.replace(old, new)

    return edit


def via_openpyxl(data: bytes, mutate: Callable[[Any], None] | None) -> bytes:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    if mutate is not None:
        mutate(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def rename_sheet(book: Any) -> None:
    book["Sheet1"].title = "Revenue"


def link_text(book: Any) -> None:
    book["Sheet1"]["A4"].value = "Expenses summary"


def header_cell(book: Any) -> None:
    book["Sheet1"]["C1"].value = "Notes"


CASES: list[dict[str, Any]] = [
    {
        "edit": "baseline (no edit)",
        "openpyxl": None,
        "targeted": {},
        "approved": [],
    },
    {
        "edit": "alt-text insertion on Picture 1 (id=2)",
        "openpyxl": "unsupported: openpyxl exposes no cNvPr title/descr API for pictures",
        "targeted": {DRAWING1: replace_once(ALT_BEFORE, ALT_AFTER)},
        "approved": [
            "drawing objects changed (count delta +0)",
            "  object id=2 fields changed: ['descr', 'title']",
        ],
    },
    {
        "edit": "worksheet rename Sheet1 -> Revenue (name only)",
        "openpyxl": rename_sheet,
        "targeted": {
            "xl/workbook.xml": replace_once(b'name="Sheet1"', b'name="Revenue"'),
        },
        "rename": {"Sheet1": "Revenue"},
        "approved": ["sheet names ['Sheet1', 'Expenses'] -> ['Revenue', 'Expenses']"],
        "check_references": True,
    },
    {
        "edit": "worksheet rename Sheet1 -> Revenue (name + formula/defined-name references)",
        "openpyxl": None,
        "targeted": {
            "xl/workbook.xml": lambda c: replace_once(b'name="Sheet1"', b'name="Revenue"')(
                c
            ).replace(b">Sheet1!", b">Revenue!"),
            "xl/worksheets/sheet2.xml": replace_once(b"<f>Sheet1!", b"<f>Revenue!"),
        },
        "rename": {"Sheet1": "Revenue"},
        "approved": [
            "sheet names ['Sheet1', 'Expenses'] -> ['Revenue', 'Expenses']",
            "Expenses!A1 formula 'Sheet1!B3*2' -> 'Revenue!B3*2'",
            "defined names {'Totals@None': 'Sheet1!$B$2:$B$3'} -> "
            "{'Totals@None': 'Revenue!$B$2:$B$3'}",
        ],
        "check_references": True,
    },
    {
        "edit": "hyperlink text A4 'Click here' -> 'Expenses summary'",
        "openpyxl": link_text,
        "targeted": {SHEET1: replace_once(b"<t>Click here</t>", b"<t>Expenses summary</t>")},
        "approved": ["Sheet1!A4 value 'Click here' -> 'Expenses summary'"],
    },
    {
        "edit": "header-row text: add C1 'Notes'",
        "openpyxl": header_cell,
        "targeted": {
            SHEET1: replace_once(
                b"<is><t>Total</t></is></c></row>",
                b'<is><t>Total</t></is></c><c r="C1" t="inlineStr"><is><t>Notes</t></is></c></row>',
            )
        },
        "approved": ["Sheet1!C1 added"],
    },
]


def stale_references(snap: dict[str, Any]) -> list[str]:
    """Formulas/defined names that reference a sheet name that no longer exists."""
    names = set(snap["sheet_order"])
    stale = []
    texts = [
        (f"{s}!{r}", c["formula"])
        for s, sh in snap["sheets"].items()
        for r, c in sh["cells"].items()
    ]
    texts += [(f"definedName {k}", v) for k, v in snap["defined_names"].items()]
    for where, text in texts:
        if text and "!" in text:
            ref = text.split("!", 1)[0].strip("'").rsplit("(", 1)[-1].rsplit("*", 1)[-1]
            if ref and ref not in names:
                stale.append(f"{where} references missing sheet {ref!r}")
    return stale


def evaluate(source: bytes, output: bytes, case: dict[str, Any]) -> dict[str, Any]:
    before = snapshot(source)
    after = snapshot(output)
    rename = case.get("rename", {})
    if rename:
        original = {new: old for old, new in rename.items()}
        after = dict(after)
        after["sheets"] = {original.get(n, n): sheet for n, sheet in after["sheets"].items()}
        after["objects"] = [
            {**obj, "worksheet": original.get(obj["worksheet"], obj["worksheet"])}
            for obj in after["objects"]
        ]
    differences = diff(before, after)
    unexplained = [d for d in differences if d not in case["approved"]]
    missing_approved = [d for d in case["approved"] if d not in differences]
    stale = stale_references(snapshot(output)) if case.get("check_references") else []
    return {
        "output_sha256": hashlib.sha256(output).hexdigest(),
        "hash_changed": hashlib.sha256(output).hexdigest() != hashlib.sha256(source).hexdigest(),
        "approved_differences_present": not missing_approved,
        "missing_approved": missing_approved,
        "unexplained_differences": unexplained,
        "stale_references": stale,
        "semantically_safe": not unexplained and not missing_approved and not stale,
        "excel_open_without_repair": "PENDING (manual check not performed)",
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--keep", type=Path, help="directory to keep output copies in")
    args = parser.parse_args()
    source = source_workbook()
    source_hash = hashlib.sha256(source).hexdigest()
    keep_dir = args.keep or Path(tempfile.mkdtemp(prefix="preservation-"))
    keep_dir.mkdir(parents=True, exist_ok=True)
    (keep_dir / "source.xlsx").write_bytes(source)
    results: list[dict[str, Any]] = []
    for index, case in enumerate(CASES):
        entry: dict[str, Any] = {"edit": case["edit"]}
        strategy = case["openpyxl"]
        if isinstance(strategy, str):
            entry["openpyxl"] = {"status": strategy}
        elif strategy is not None or index == 0:
            output = via_openpyxl(source, strategy)
            (keep_dir / f"case{index}_openpyxl.xlsx").write_bytes(output)
            entry["openpyxl"] = evaluate(source, output, case)
        if case["targeted"] or index == 0:
            output = targeted(source, case["targeted"])
            (keep_dir / f"case{index}_targeted.xlsx").write_bytes(output)
            entry["targeted"] = evaluate(source, output, case)
        results.append(entry)
    assert hashlib.sha256(source).hexdigest() == source_hash
    report = {
        "source_sha256": source_hash,
        "source_unchanged": True,
        "outputs_directory": "temporary (not committed)",
        "results": results,
    }
    print(json.dumps(report, indent=2, ensure_ascii=False))
    print(f"\nOutput copies written to {keep_dir}", file=sys.stderr)


if __name__ == "__main__":
    main()
