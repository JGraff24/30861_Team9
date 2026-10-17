"""Dylan: safe intake, rejection codes, relationships and nonempty-cell counting."""

import hashlib
import zipfile

import pytest

from app.audit.engine import OoxmlWorkbookEngine
from app.contracts.errors import WorkbookValidationError
from app.contracts.models import ProcessingLimits
from tests.unit.ooxml_builder import (
    MAIN,
    REL,
    Sheet,
    Workbook,
    rels_xml,
    set_encrypted_flag,
    simple_workbook,
    zip_parts,
)

ENGINE = OoxmlWorkbookEngine()
LIMITS = ProcessingLimits()


def validate(data: bytes, filename: str = "book.xlsx", limits: ProcessingLimits = LIMITS):  # type: ignore[no-untyped-def]
    return ENGINE.validate(data, filename, limits)


def rejection(data: bytes, filename: str = "book.xlsx") -> WorkbookValidationError:
    before = hashlib.sha256(data).hexdigest()
    with pytest.raises(WorkbookValidationError) as caught:
        validate(data, filename)
    assert hashlib.sha256(data).hexdigest() == before
    return caught.value


def replace_part(wb: Workbook, name: str, content: bytes | None) -> bytes:
    parts = wb.parts()
    if content is None:
        parts.pop(name)
    else:
        parts[name] = content
    return zip_parts(parts)


# -- accepted inputs -------------------------------------------------------------------


def test_valid_workbook_returns_exact_hash_size_counts_and_part_map() -> None:
    wb = Workbook(
        [
            Sheet(
                "Expenses",
                "xl/worksheets/a_first.xml",
                "4",
                cells='<row r="1"><c r="A1"><v>5</v></c></row>',
            ),
            Sheet("Sheet1", "xl/worksheets/zzz.xml", "9", state="hidden"),
            Sheet(
                "Archive",
                "xl/worksheets/sheet1.xml",
                "2",
                state="veryHidden",
                cells='<row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row>',
            ),
        ]
    )
    data = wb.build()
    result = validate(data, "Q1 BUDGET.XLSX")
    assert result.source_sha256 == hashlib.sha256(data).hexdigest()
    assert result.size_bytes == len(data)
    assert result.worksheet_count == 3
    assert result.nonempty_cell_count == 3
    # Names do not match filenames; mapping follows the actual relationships.
    assert result.worksheet_parts == {
        "Expenses": "xl/worksheets/a_first.xml",
        "Sheet1": "xl/worksheets/zzz.xml",
        "Archive": "xl/worksheets/sheet1.xml",
    }


def test_nonempty_cell_convention() -> None:
    cells = (
        '<row r="1">'
        '<c r="A1"><v>0</v></c>'  # numeric zero: counts
        '<c r="B1" t="b"><v>0</v></c>'  # boolean false: counts
        '<c r="C1"><f>A1+1</f></c>'  # formula without cached value: counts
        '<c r="D1"><f>A1+2</f><v>2</v></c>'  # formula with value: counts once
        '<c r="E1" t="s"><v>0</v></c>'  # shared string reference: counts
        '<c r="F1" t="inlineStr"><is><t>Hi</t></is></c>'  # inline text: counts
        '<c r="G1" t="inlineStr"><is><r><t>rich</t></r></is></c>'  # rich inline: counts
        '<c r="H1" t="str"><v> </v></c>'  # whitespace value is not empty: counts
        '<c r="I1" s="3"/>'  # formatting only: no
        '<c r="J1"><v></v></c>'  # empty value: no
        '<c r="K1" t="inlineStr"><is><t></t></is></c>'  # empty inline string: no
        '<c r="L1" t="inlineStr"><is/></c>'  # empty inline container: no
        '</row><row r="2" s="1" customFormat="1"/>'
    )
    data = simple_workbook("Data", cells=cells).build()
    assert validate(data).nonempty_cell_count == 8


def test_cells_on_hidden_sheets_count() -> None:
    wb = simple_workbook("Visible", "Hidden")
    wb.sheets[1].state = "hidden"
    assert validate(wb.build()).nonempty_cell_count == 2


def test_relationship_resolution_preserves_stored_case_and_handles_absolute_targets() -> None:
    wb = simple_workbook("Totals", "Notes")
    parts = wb.parts()
    parts["xl/Worksheets/Sheet2.XML"] = parts.pop("xl/worksheets/sheet2.xml")
    parts["xl/_rels/workbook.xml.rels"] = rels_xml(
        f'<Relationship Id="rIdSheet1" Type="{REL}/worksheet" Target="/xl/worksheets/sheet1.xml"/>'
        f'<Relationship Id="rIdSheet2" Type="{REL}/worksheet" '
        'Target="./worksheets/../worksheets/sheet2.xml"/>'
    )
    ct = (
        parts["[Content_Types].xml"]
        .decode()
        .replace("/xl/worksheets/sheet2.xml", "/xl/Worksheets/Sheet2.XML")
    )
    parts["[Content_Types].xml"] = ct.encode()
    result = validate(zip_parts(parts))
    assert result.worksheet_parts == {
        "Totals": "xl/worksheets/sheet1.xml",
        "Notes": "xl/Worksheets/Sheet2.XML",
    }


def test_strict_ooxml_namespaces_are_accepted() -> None:
    wb = simple_workbook("Data")
    parts = wb.parts()
    strict_main = "http://purl.oclc.org/ooxml/spreadsheetml/main"
    strict_rel = "http://purl.oclc.org/ooxml/officeDocument/relationships"
    for name in ("xl/workbook.xml", "xl/worksheets/sheet1.xml"):
        parts[name] = (
            parts[name]
            .replace(MAIN.encode(), strict_main.encode())
            .replace(REL.encode(), strict_rel.encode())
        )
    for name in ("_rels/.rels", "xl/_rels/workbook.xml.rels"):
        parts[name] = parts[name].replace(REL.encode(), strict_rel.encode())
    assert validate(zip_parts(parts)).worksheet_parts == {"Data": "xl/worksheets/sheet1.xml"}


def test_stored_compression_and_uppercase_extension_accepted() -> None:
    wb = simple_workbook("Data")
    wb.compression = zipfile.ZIP_STORED
    assert validate(wb.build(), "C:\\Users\\me\\Report.XLSX").worksheet_count == 1


def test_openpyxl_written_workbook_is_accepted(tmp_path) -> None:  # type: ignore[no-untyped-def]
    from openpyxl import Workbook as OpenpyxlWorkbook

    book = OpenpyxlWorkbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet["A1"] = 0
    sheet["A2"] = False
    sheet["A3"] = "=A1+1"
    sheet["A4"] = "text"
    hidden = book.create_sheet("Lookup")
    hidden.sheet_state = "hidden"
    hidden["B2"] = 3.5
    path = tmp_path / "written.xlsx"
    book.save(path)
    result = validate(path.read_bytes())
    assert result.worksheet_count == 2
    assert result.nonempty_cell_count == 5
    assert set(result.worksheet_parts) == {"Budget", "Lookup"}


# -- file type -------------------------------------------------------------------------


@pytest.mark.parametrize(
    "filename", ["book.xls", "book.csv", "book.txt", "book", "", "book.xlsx.zip", ".xlsx/evil"]
)
def test_unsupported_extensions_are_415(filename: str) -> None:
    error = rejection(simple_workbook("Data").build(), filename)
    assert (error.status_code, error.code) == (415, "UNSUPPORTED_FILE_TYPE")


def test_named_xlsm_is_422_macro() -> None:
    error = rejection(simple_workbook("Data").build(), "book.XLSM")
    assert (error.status_code, error.code) == (422, "MACRO_ENABLED_WORKBOOK_UNSUPPORTED")


# -- macros ----------------------------------------------------------------------------


def test_xlsx_with_vba_part_is_macro_rejected() -> None:
    wb = simple_workbook("Data")
    wb.extra_parts["xl/vbaProject.bin"] = b"\x00binary"
    assert rejection(wb.build()).code == "MACRO_ENABLED_WORKBOOK_UNSUPPORTED"


def test_macro_enabled_main_content_type_is_rejected() -> None:
    wb = simple_workbook("Data")
    wb.workbook_content_type = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
    assert rejection(wb.build()).code == "MACRO_ENABLED_WORKBOOK_UNSUPPORTED"


def test_vba_relationship_is_rejected_even_without_bin_part() -> None:
    wb = simple_workbook("Data")
    parts = wb.parts()
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/'
        b'relationships/vbaProject" Target="code.dat"/></Relationships>',
    )
    assert rejection(zip_parts(parts)).code == "MACRO_ENABLED_WORKBOOK_UNSUPPORTED"


def test_excel4_macro_sheet_part_is_rejected() -> None:
    wb = simple_workbook("Data")
    wb.extra_parts["xl/macrosheets/sheet1.xml"] = b"<xm/>"
    assert rejection(wb.build()).code == "MACRO_ENABLED_WORKBOOK_UNSUPPORTED"


# -- corrupt, encrypted and malformed --------------------------------------------------


@pytest.mark.parametrize(
    "data",
    [b"", b"not a zip at all", b"PK\x03\x04truncated", b"col1,col2\n1,2\n"],
    ids=["empty", "text", "truncated-zip", "csv-content"],
)
def test_unreadable_bytes_are_invalid(data: bytes) -> None:
    error = rejection(data)
    assert (error.status_code, error.code) == (422, "INVALID_WORKBOOK")


def test_truncated_real_workbook_is_invalid() -> None:
    data = simple_workbook("Data").build()
    assert rejection(data[: len(data) // 2]).code == "INVALID_WORKBOOK"


def test_corrupted_entry_bytes_are_invalid() -> None:
    wb = simple_workbook("Data")
    wb.compression = zipfile.ZIP_STORED
    data = bytearray(wb.build())
    marker = data.find(b"<sheetData>")
    data[marker + 1] ^= 0xFF  # CRC no longer matches
    assert rejection(bytes(data)).code == "INVALID_WORKBOOK"


def test_ole_encrypted_package_is_encrypted_rejection() -> None:
    data = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 64 + "EncryptionInfo".encode("utf-16-le")
    error = rejection(data)
    assert (error.status_code, error.code) == (422, "ENCRYPTED_WORKBOOK_UNSUPPORTED")


def test_legacy_binary_renamed_xlsx_is_invalid() -> None:
    assert rejection(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512).code == "INVALID_WORKBOOK"


def test_encrypted_zip_entry_is_rejected() -> None:
    data = set_encrypted_flag(simple_workbook("Data").build(), "xl/worksheets/sheet1.xml")
    assert rejection(data).code == "ENCRYPTED_WORKBOOK_UNSUPPORTED"


@pytest.mark.parametrize(
    "missing",
    [
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/worksheets/sheet1.xml",
    ],
)
def test_missing_required_parts_are_invalid(missing: str) -> None:
    error = rejection(replace_part(simple_workbook("Data"), missing, None))
    assert error.code == "INVALID_WORKBOOK"


@pytest.mark.parametrize(
    "part",
    [
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
        "[Content_Types].xml",
        "xl/_rels/workbook.xml.rels",
    ],
)
def test_malformed_xml_is_invalid(part: str) -> None:
    truncated = b'<worksheet xmlns="' + MAIN.encode() + b'"><sheetData><row r="1">'
    error = rejection(replace_part(simple_workbook("Data"), part, truncated))
    assert (error.code, error.details) == ("INVALID_WORKBOOK", {"reason": "malformed_xml"})


def test_malformed_unreferenced_xml_part_is_invalid() -> None:
    wb = simple_workbook("Data")
    wb.extra_parts["docProps/custom.xml"] = b"<a><b></a>"
    assert rejection(wb.build()).details == {"reason": "malformed_xml"}


def test_wrong_workbook_content_type_is_invalid() -> None:
    wb = simple_workbook("Data")
    wb.workbook_content_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    )
    assert rejection(wb.build()).details == {"reason": "workbook_content_type"}


def test_zero_sheets_is_invalid() -> None:
    assert rejection(Workbook([]).build()).details == {"reason": "no_sheets"}


def test_duplicate_sheet_names_are_invalid() -> None:
    wb = simple_workbook("Data", "DATA")
    assert rejection(wb.build()).details == {"reason": "duplicate_sheet"}


def test_external_worksheet_target_is_invalid_and_not_fetched() -> None:
    wb = simple_workbook("Data")
    parts = wb.parts()
    parts["xl/_rels/workbook.xml.rels"] = rels_xml(
        f'<Relationship Id="rIdSheet1" Type="{REL}/worksheet" '
        'Target="http://127.0.0.1:9/sheet.xml" TargetMode="External"/>'
    )
    assert rejection(zip_parts(parts)).details == {"reason": "missing_sheet_part"}


# -- XML attacks and unsafe package paths ----------------------------------------------

BILLION_LAUGHS = (
    b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
    b'<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
    b'<worksheet xmlns="' + MAIN.encode() + b'"><sheetData/><x>&lol2;</x></worksheet>'
)
XXE = (
    b'<?xml version="1.0"?><!DOCTYPE w [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
    b'<worksheet xmlns="' + MAIN.encode() + b'"><sheetData/><x>&xxe;</x></worksheet>'
)
BARE_DTD = (
    b'<?xml version="1.0"?><!DOCTYPE worksheet><worksheet xmlns="'
    + MAIN.encode()
    + b'"><sheetData/></worksheet>'
)


@pytest.mark.parametrize("payload", [BILLION_LAUGHS, XXE, BARE_DTD], ids=["entities", "xxe", "dtd"])
def test_dtd_and_entities_are_unsafe(payload: bytes) -> None:
    error = rejection(replace_part(simple_workbook("Data"), "xl/worksheets/sheet1.xml", payload))
    assert (error.code, error.details) == (
        "UNSAFE_WORKBOOK_PACKAGE",
        {"reason": "xml_dtd_or_entity"},
    )


def test_dtd_in_unreferenced_part_is_unsafe() -> None:
    wb = simple_workbook("Data")
    wb.extra_parts["customXml/item1.xml"] = BARE_DTD
    assert rejection(wb.build()).code == "UNSAFE_WORKBOOK_PACKAGE"


@pytest.mark.parametrize(
    "entry",
    [
        "../evil.xml",
        "/abs.xml",
        "xl\\win.xml",
        "xl/./x.xml",
        "xl//x.xml",
        "C:/x.xml",
        "xl/\x01.xml",
    ],
)
def test_unsafe_entry_paths_are_rejected(entry: str) -> None:
    parts = list(simple_workbook("Data").parts().items()) + [(entry, b"<x/>")]
    error = rejection(zip_parts(parts))
    assert (error.code, error.details) == ("UNSAFE_WORKBOOK_PACKAGE", {"reason": "entry_path"})


@pytest.mark.parametrize("duplicate", ["xl/workbook.xml", "XL/Workbook.xml"])
def test_duplicate_entries_are_rejected(duplicate: str) -> None:
    parts = list(simple_workbook("Data").parts().items()) + [(duplicate, b"<x/>")]
    assert rejection(zip_parts(parts)).details == {"reason": "duplicate_entry"}


@pytest.mark.parametrize(
    "target", ["../../outside.xml", "/../outside.xml", "worksheets/../../../x.xml"]
)
def test_relationship_traversal_is_rejected(target: str) -> None:
    wb = simple_workbook("Data")
    parts = wb.parts()
    parts["xl/_rels/workbook.xml.rels"] = rels_xml(
        f'<Relationship Id="rIdSheet1" Type="{REL}/worksheet" Target="{target}"/>'
    )
    assert rejection(zip_parts(parts)).details == {"reason": "relationship_target"}


def test_traversal_in_any_rels_part_is_rejected() -> None:
    wb = simple_workbook("Data")
    wb.extra_parts["xl/worksheets/_rels/sheet1.xml.rels"] = rels_xml(
        f'<Relationship Id="rId1" Type="{REL}/image" Target="../../../../etc/passwd"/>'
    )
    assert rejection(wb.build()).code == "UNSAFE_WORKBOOK_PACKAGE"


def test_error_details_never_echo_entry_names() -> None:
    parts = list(simple_workbook("Data").parts().items()) + [("../secret-name.xml", b"<x/>")]
    error = rejection(zip_parts(parts))
    assert "secret" not in repr(error.details) and "secret" not in error.message


# -- protection policy -----------------------------------------------------------------


@pytest.mark.parametrize(
    "head",
    [
        '<workbookProtection lockStructure="1"/>',
        '<workbookProtection lockWindows="true"/>',
        '<workbookProtection workbookAlgorithmName="SHA-512" workbookHashValue="x" '
        'workbookSaltValue="y"/>',
        '<workbookProtection workbookPassword="CC1A"/>',
        '<fileSharing userName="a" reservationPassword="CC1A"/>',
    ],
)
def test_protected_workbook_structures_are_rejected(head: str) -> None:
    wb = simple_workbook("Data")
    wb.workbook_head = head
    error = rejection(wb.build())
    assert (error.code, error.details) == ("PROTECTED_WORKBOOK_UNSUPPORTED", {"scope": "workbook"})


@pytest.mark.parametrize(
    "head", ['<workbookProtection lockStructure="0"/>', '<fileSharing readOnlyRecommended="1"/>']
)
def test_unlocked_protection_elements_are_accepted(head: str) -> None:
    wb = simple_workbook("Data")
    wb.workbook_head = head
    assert validate(wb.build()).worksheet_count == 1


@pytest.mark.parametrize(
    "tail",
    [
        '<sheetProtection sheet="1" objects="1"/>',
        '<sheetProtection sheet="true"/>',
        '<sheetProtection password="CC1A" sheet="0"/>',
        '<sheetProtection algorithmName="SHA-512" hashValue="x" saltValue="y" spinCount="1"/>',
    ],
)
def test_protected_worksheets_are_rejected(tail: str) -> None:
    wb = simple_workbook("Data", "Hidden")
    wb.sheets[1].state = "hidden"
    wb.sheets[1].tail = tail
    assert rejection(wb.build()).details == {"scope": "worksheet"}


def test_sheet_protection_off_is_accepted() -> None:
    wb = simple_workbook("Data")
    wb.sheets[0].tail = '<sheetProtection sheet="0" objects="1"/>'
    assert validate(wb.build()).worksheet_count == 1
